"""doctool writer のテスト (要 python-docx)。

実行: PYTHONUTF8=1 python .claude/skills/doctool/scripts/doctool/tests/test_docx_writer.py
"""
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import docx_reader

try:
    import docx_writer
    HAS_PYTHON_DOCX = True
except ImportError:
    HAS_PYTHON_DOCX = False

REPO_ROOT = Path(__file__).resolve().parents[6]
HANREI_DIR = REPO_ROOT / 'docs' / 'specs' / '詳細設計_書式凡例'
TEMPLATE = HANREI_DIR / '型情報定義書' / '型情報定義書一覧表_凡例.docx'
# 複数の表が同一 numId を共有している凡例(2 つの採番表が numId=30 を共有)。
NUM_TEMPLATE = HANREI_DIR / 'イベント定義書' / 'コントローラー一覧表_凡例.docx'


def _numids_per_table(path):
    """body 直下の各表が参照する numId 集合のリスト(採番のある表のみ)。"""
    from docx import Document
    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    doc = Document(path)
    out = []
    for tbl in doc.element.body.iter(W + 'tbl'):
        ids = {nid.get(W + 'val') for nid in tbl.iter(W + 'numId')}
        ids.discard(None)
        if ids:
            out.append(ids)
    return out


@unittest.skipUnless(HAS_PYTHON_DOCX, 'python-docx なし')
@unittest.skipUnless(TEMPLATE.is_file(), '凡例なし')
class WriterEndToEndTest(unittest.TestCase):
    """テーブル定義書一覧表の生成(過去に手書きスクリプトで行った内容)をスペックで再現する。"""

    @classmethod
    def setUpClass(cls):
        # 雛形の IR から対象ブロックの座標を特定する(想定する利用フロー)
        ir = docx_reader.read_docx(TEMPLATE, source_name=TEMPLATE.name)
        cls.hdr_blocks = ir['sections'][0]['headers']['default']
        cls.hdr_table = next(i for i, b in enumerate(cls.hdr_blocks)
                             if b['type'] == 'table')
        body = ir['body']
        # 本文構造: P(基本型情報)/表 / P(モデル型定義一覧表…)/表 / P(列挙型定義一覧表…)/表 / P(空)
        # 2 つ目の一覧表(モデル型)をテーブル定義一覧表に作り直すシナリオで生成機構を検証する
        cls.title_block = next(i for i, b in enumerate(body)
                               if b.get('text', '').startswith('モデル型定義一覧表'))
        cls.list_table = cls.title_block + 1

        cls.tmpdir = tempfile.TemporaryDirectory()
        cls.out = str(Path(cls.tmpdir.name) / 'out.docx')
        cls.warnings = docx_writer.apply_spec({
            'template': str(TEMPLATE),
            'output': cls.out,
            'operations': [
                {'op': 'setCell', 'part': 'header', 'block': cls.hdr_table,
                 'row': 0, 'col': 0, 'text': 'テーブル定義書一覧表'},
                {'op': 'setCell', 'part': 'header', 'block': cls.hdr_table,
                 'row': 0, 'col': 5, 'text': '2026-06-10'},
                {'op': 'setCell', 'part': 'header', 'block': cls.hdr_table,
                 'row': 1, 'col': 3, 'text': 'テーブル定義一覧表'},
                {'op': 'deleteBlock', 'part': 'body', 'block': 0},
                {'op': 'deleteBlock', 'part': 'body', 'block': 1},
                {'op': 'setParagraph', 'part': 'body', 'block': cls.title_block,
                 'text': 'テーブル定義一覧表'},
                {'op': 'fillTable', 'part': 'body', 'block': cls.list_table,
                 'dataRow': 1, 'rows': [
                     [None, 'アカウントテーブル', 'accounts', 'アカウント情報を格納するテーブル'],
                     [None, '注文テーブル', 'orders', '注文情報を格納するテーブル'],
                 ]},
            ],
        })
        cls.result = docx_reader.read_docx(cls.out, source_name='out.docx')

    @classmethod
    def tearDownClass(cls):
        cls.tmpdir.cleanup()

    def test_no_warnings(self):
        self.assertEqual(self.warnings, [])

    def test_header_replaced(self):
        header = self.result['sections'][0]['headers']['default']
        table = next(b for b in header if b['type'] == 'table')
        cell = {(c['row'], c['col']): c for c in table['cells']}
        self.assertEqual(cell[(0, 0)]['blocks'][0]['text'], 'テーブル定義書一覧表')
        self.assertEqual(cell[(0, 5)]['blocks'][0]['text'], '2026-06-10')
        self.assertEqual(cell[(1, 3)]['blocks'][0]['text'], 'テーブル定義一覧表')

    def test_body_structure(self):
        body = self.result['body']
        # 基本型情報の見出しと表が消え、先頭がタイトル段落になる
        self.assertEqual(body[0]['text'], 'テーブル定義一覧表')
        table = body[1]
        self.assertEqual(table['type'], 'table')
        self.assertEqual(table['rowCount'], 3)  # ヘッダー行 + データ2行
        cell = {(c['row'], c['col']): c for c in table['cells']}
        self.assertEqual(cell[(1, 1)]['blocks'][0]['text'], 'アカウントテーブル')
        self.assertEqual(cell[(2, 2)]['blocks'][0]['text'], 'orders')

    def test_auto_number_preserved(self):
        table = self.result['body'][1]
        for row in (1, 2):
            no_cell = next(c for c in table['cells']
                           if c['row'] == row and c['col'] == 0)
            p = no_cell['blocks'][0]
            self.assertIn('numbering', p, f'行{row}のNo列に採番がない')
            self.assertEqual(p['text'], '')


@unittest.skipUnless(HAS_PYTHON_DOCX, 'python-docx なし')
@unittest.skipUnless(TEMPLATE.is_file(), '凡例なし')
class WriterGuardTest(unittest.TestCase):
    def _spec(self, operations, out):
        return {'template': str(TEMPLATE), 'output': out, 'operations': operations}

    def test_numbering_guard_warns(self):
        ir = docx_reader.read_docx(TEMPLATE, source_name=TEMPLATE.name)
        list_table = next(
            i for i, b in enumerate(ir['body'])
            if i > 0 and b['type'] == 'table')
        with tempfile.TemporaryDirectory() as tmp:
            out = str(Path(tmp) / 'out.docx')
            warnings = docx_writer.apply_spec(self._spec([
                {'op': 'fillTable', 'part': 'body', 'block': list_table,
                 'dataRow': 1, 'rows': [['1', 'X', 'Y', 'Z']]},
            ], out))
            self.assertTrue(any('自動採番' in w for w in warnings))
            # No列は雛形のまま(空+採番)、他の列は設定される
            table = next(b for b in docx_reader.read_docx(out)['body']
                         if b['type'] == 'table')
            cell = {(c['row'], c['col']): c for c in table['cells']}
            self.assertEqual(cell[(1, 0)]['blocks'][0]['text'], '')
            self.assertEqual(cell[(1, 1)]['blocks'][0]['text'], 'X')

    def test_append_section_clones_blocks_with_format(self):
        # 見出し(基本型情報)と表を複製して末尾へ節を追加する
        with tempfile.TemporaryDirectory() as tmp:
            out = str(Path(tmp) / 'out.docx')
            warnings = docx_writer.apply_spec(self._spec([
                {'op': 'appendSection', 'part': 'body', 'blocks': [
                    {'cloneFrom': 0, 'text': '追加の節'},
                    {'cloneFrom': 1, 'dataRow': 1, 'rows': [
                        [None, '論理名X', 'physicalX', '概要X'],
                    ]},
                ]},
            ], out))
            self.assertEqual(warnings, [])
            ir = docx_reader.read_docx(out)
            body = ir['body']
            # 末尾(元の最終ブロックの後)に P+TBL が追加されている
            self.assertEqual(body[-2]['text'], '追加の節')
            added = body[-1]
            self.assertEqual(added['type'], 'table')
            self.assertEqual(added['rowCount'], 2)
            cell = {(c['row'], c['col']): c for c in added['cells']}
            self.assertEqual(cell[(1, 1)]['blocks'][0]['text'], '論理名X')
            # 複製元 (基本型情報の表) は変更されていない
            first_table = next(b for b in body if b['type'] == 'table')
            self.assertEqual(first_table['rowCount'], 8)

    def test_covered_cell_rejected(self):
        ir = docx_reader.read_docx(TEMPLATE, source_name=TEMPLATE.name)
        hdr_table = next(i for i, b in enumerate(
            ir['sections'][0]['headers']['default']) if b['type'] == 'table')
        # ヘッダー左上は縦結合(rowSpan2)のため (1,0) は被覆側 → エラー
        with self.assertRaises(ValueError):
            docx_writer.apply_spec(self._spec([
                {'op': 'setCell', 'part': 'header', 'block': hdr_table,
                 'row': 1, 'col': 0, 'text': 'NG'},
            ], 'unused.docx'))


@unittest.skipUnless(HAS_PYTHON_DOCX, 'python-docx なし')
class ReplaceTextTest(unittest.TestCase):
    """replaceText: ラン書式を保持した部分置換(レビュー指摘の修正用)。"""

    def _build(self, tmp):
        from docx import Document
        doc = Document()
        p = doc.add_paragraph()
        p.add_run('重要な').bold = True
        p.add_run('内容を誤入力した。誤りは直す。')
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'セルの誤記'
        src = str(Path(tmp) / 'src.docx')
        doc.save(src)
        return src

    def _apply(self, tmp, operations):
        out = str(Path(tmp) / 'out.docx')
        warnings = docx_writer.apply_spec({
            'template': self._build(tmp), 'output': out,
            'operations': operations,
        })
        return docx_reader.read_docx(out, source_name='out.docx'), warnings

    def test_replace_first_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            ir, warnings = self._apply(tmp, [
                {'op': 'replaceText', 'part': 'body', 'block': 0,
                 'find': '誤', 'replace': '正'},
            ])
            self.assertEqual(warnings, [])
            self.assertEqual(ir['body'][0]['text'], '重要な内容を正入力した。誤りは直す。')

    def test_replace_all(self):
        with tempfile.TemporaryDirectory() as tmp:
            ir, warnings = self._apply(tmp, [
                {'op': 'replaceText', 'part': 'body', 'block': 0,
                 'find': '誤', 'replace': '正', 'all': True},
            ])
            self.assertEqual(warnings, [])
            self.assertEqual(ir['body'][0]['text'], '重要な内容を正入力した。正りは直す。')

    def test_cross_run_match_keeps_first_run_format(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = str(Path(tmp) / 'out.docx')
            warnings = docx_writer.apply_spec({
                'template': self._build(tmp), 'output': out,
                'operations': [
                    # 「重要な内」は太字ラン+通常ランにまたがる
                    {'op': 'replaceText', 'part': 'body', 'block': 0,
                     'find': '重要な内', 'replace': '大事な内'},
                ],
            })
            self.assertEqual(warnings, [])
            ir = docx_reader.read_docx(out, with_format=True)
            p = ir['body'][0]
            self.assertEqual(p['text'], '大事な内容を誤入力した。誤りは直す。')
            self.assertEqual(p['runs'][0], {'text': '大事な内', 'bold': True})

    def test_replace_in_cell(self):
        with tempfile.TemporaryDirectory() as tmp:
            ir, warnings = self._apply(tmp, [
                {'op': 'replaceText', 'part': 'body', 'block': 1,
                 'row': 0, 'col': 0, 'find': '誤記', 'replace': '記述'},
            ])
            self.assertEqual(warnings, [])
            table = next(b for b in ir['body'] if b['type'] == 'table')
            self.assertEqual(table['cells'][0]['blocks'][0]['text'], 'セルの記述')

    def test_not_found_warns(self):
        with tempfile.TemporaryDirectory() as tmp:
            _, warnings = self._apply(tmp, [
                {'op': 'replaceText', 'part': 'body', 'block': 0,
                 'find': '存在しない語', 'replace': 'X'},
            ])
            self.assertTrue(any('見つかりません' in w for w in warnings))

    def test_table_without_cell_address_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                self._apply(tmp, [
                    {'op': 'replaceText', 'part': 'body', 'block': 1,
                     'find': '誤記', 'replace': 'X'},
                ])


@unittest.skipUnless(HAS_PYTHON_DOCX, 'python-docx なし')
class MixedFormatSetCellTest(unittest.TestCase):
    """setCell: 混在書式段落では変化した中間部分のみ差し替えて書式を維持する。"""

    def _set(self, tmp, text):
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        p = table.cell(0, 0).paragraphs[0]
        p.add_run('種別: ').bold = True
        p.add_run('未記入')
        src = str(Path(tmp) / 'src.docx')
        doc.save(src)
        out = str(Path(tmp) / 'out.docx')
        warnings = docx_writer.apply_spec({
            'template': src, 'output': out, 'operations': [
                {'op': 'setCell', 'part': 'body', 'block': 0,
                 'row': 0, 'col': 0, 'text': text},
            ],
        })
        ir = docx_reader.read_docx(out, with_format=True)
        return ir['body'][0]['cells'][0]['blocks'][0], warnings

    def test_value_change_keeps_label_format(self):
        with tempfile.TemporaryDirectory() as tmp:
            p, warnings = self._set(tmp, '種別: 機能要件')
            self.assertEqual(warnings, [])
            self.assertEqual(p['text'], '種別: 機能要件')
            self.assertEqual(p['runs'], [
                {'text': '種別: ', 'bold': True},
                {'text': '機能要件'},
            ])

    def test_full_rewrite_warns_and_flattens(self):
        with tempfile.TemporaryDirectory() as tmp:
            p, warnings = self._set(tmp, '全く別の内容')
            self.assertEqual(p['text'], '全く別の内容')
            self.assertTrue(any('均され' in w for w in warnings))


@unittest.skipUnless(HAS_PYTHON_DOCX, 'python-docx なし')
class SetFormatTest(unittest.TestCase):
    """setFormat: ラン・段落・セルへの書式設定(レビュー指摘の書式修正用)。"""

    def _build(self, tmp):
        from docx import Document
        doc = Document()
        p = doc.add_paragraph()
        p.add_run('重要な').bold = True
        p.add_run('内容を誤入力した。')
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'セル見出し'
        src = str(Path(tmp) / 'src.docx')
        doc.save(src)
        return src

    def _apply(self, tmp, operations):
        out = str(Path(tmp) / 'out.docx')
        warnings = docx_writer.apply_spec({
            'template': self._build(tmp), 'output': out,
            'operations': operations,
        })
        return docx_reader.read_docx(out, with_format=True), warnings

    def test_whole_paragraph_color(self):
        with tempfile.TemporaryDirectory() as tmp:
            ir, warnings = self._apply(tmp, [
                {'op': 'setFormat', 'part': 'body', 'block': 0,
                 'format': {'color': 'FF0000'}},
            ])
            self.assertEqual(warnings, [])
            self.assertEqual(ir['body'][0]['runs'], [
                {'text': '重要な', 'bold': True, 'color': 'FF0000'},
                {'text': '内容を誤入力した。', 'color': 'FF0000'},
            ])

    def test_span_format_splits_runs(self):
        with tempfile.TemporaryDirectory() as tmp:
            ir, warnings = self._apply(tmp, [
                {'op': 'setFormat', 'part': 'body', 'block': 0,
                 'find': '誤入力', 'format': {'bold': True, 'color': 'FF0000'}},
            ])
            self.assertEqual(warnings, [])
            p = ir['body'][0]
            self.assertEqual(p['text'], '重要な内容を誤入力した。')
            self.assertEqual(p['runs'], [
                {'text': '重要な', 'bold': True},
                {'text': '内容を'},
                {'text': '誤入力', 'bold': True, 'color': 'FF0000'},
                {'text': 'した。'},
            ])

    def test_bold_off_overrides(self):
        with tempfile.TemporaryDirectory() as tmp:
            ir, warnings = self._apply(tmp, [
                {'op': 'setFormat', 'part': 'body', 'block': 0,
                 'format': {'bold': False}},
            ])
            self.assertEqual(warnings, [])
            # w:b w:val="0" の明示上書き → 読み戻しでは書式なし扱い
            self.assertNotIn('runs', ir['body'][0])

    def test_cell_align_and_shading(self):
        with tempfile.TemporaryDirectory() as tmp:
            ir, warnings = self._apply(tmp, [
                {'op': 'setFormat', 'part': 'body', 'block': 1,
                 'row': 0, 'col': 0,
                 'format': {'align': 'center', 'shading': 'FFFF00'}},
            ])
            self.assertEqual(warnings, [])
            cell = next(b for b in ir['body'] if b['type'] == 'table')['cells'][0]
            self.assertEqual(cell['shading'], 'FFFF00')
            self.assertEqual(cell['blocks'][0]['align'], 'center')

    def test_find_not_found_warns(self):
        with tempfile.TemporaryDirectory() as tmp:
            _, warnings = self._apply(tmp, [
                {'op': 'setFormat', 'part': 'body', 'block': 0,
                 'find': '存在しない語', 'format': {'bold': True}},
            ])
            self.assertTrue(any('見つかりません' in w for w in warnings))

    def test_unknown_format_key_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                self._apply(tmp, [
                    {'op': 'setFormat', 'part': 'body', 'block': 0,
                     'format': {'fontSize': 12}},
                ])


@unittest.skipUnless(HAS_PYTHON_DOCX, 'python-docx なし')
@unittest.skipUnless(NUM_TEMPLATE.is_file(), '採番凡例なし')
class TableNumberingNormalizeTest(unittest.TestCase):
    """表ごとに独立した自動採番(表をまたいだ連番継続の防止)。"""

    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

    def _doc_with_shared_numid(self):
        """採番表を複製し、2 表が同一 numId を共有する(連番が地続きになる)
        文書を組み立てて返す。凡例ファイルの現状に依存しないテスト用の素地。"""
        import copy
        from docx import Document
        W = self.W
        doc = Document(str(NUM_TEMPLATE))
        body = doc.element.body
        src = next(t for t in body.findall(W + 'tbl')
                   if t.find('.//' + W + 'numPr') is not None)
        clone = copy.deepcopy(src)
        sect = body.find(W + 'sectPr')
        (sect.addprevious if sect is not None else body.append)(clone)
        return doc, src, clone

    def _ids(self, tbl):
        W = self.W
        return {n.get(W + 'val') for n in tbl.iter(W + 'numId')}

    def test_shared_numid_is_split_per_table(self):
        doc, src, clone = self._doc_with_shared_numid()
        before = self._ids(src)
        self.assertTrue(before & self._ids(clone))  # 複製直後は numId 共有
        created = docx_writer.normalize_table_numbering(doc)
        self.assertGreaterEqual(created, 1)
        self.assertEqual(self._ids(src), before)            # 最初の表は元のまま
        self.assertTrue(self._ids(src).isdisjoint(self._ids(clone)))  # 複製は別採番
        # 冪等: 2 回目は新規割当てなし
        self.assertEqual(docx_writer.normalize_table_numbering(doc), 0)

    def test_cloned_numbered_table_gets_own_numid(self):
        # 採番表を複製しても、複製先は独立採番(1 始まり)になる
        ir = docx_reader.read_docx(NUM_TEMPLATE)
        tbl_block = next(i for i, b in enumerate(ir['body'])
                         if b['type'] == 'table')
        with tempfile.TemporaryDirectory() as tmp:
            out = str(Path(tmp) / 'out.docx')
            docx_writer.apply_spec({
                'template': str(NUM_TEMPLATE), 'output': out, 'operations': [
                    {'op': 'appendSection', 'part': 'body', 'blocks': [
                        {'cloneFrom': tbl_block, 'dataRow': 1,
                         'rows': [[None, 'X', 'Y', 'Z', 'W']]},
                    ]},
                ],
            })
            ids = _numids_per_table(out)
            # 元の 2 表 + 複製 1 表 = 3 表、すべて相異なる numId
            flat = [x for s in ids for x in s]
            self.assertEqual(len(flat), len(set(flat)), f'numId 重複: {ids}')

    def test_normalize_file_in_place(self):
        # apply_spec を通さず、numId 共有のまま保存したファイルを是正する
        doc, _, _ = self._doc_with_shared_numid()
        with tempfile.TemporaryDirectory() as tmp:
            target = str(Path(tmp) / 'shared.docx')
            doc.save(target)
            self.assertTrue(docx_writer.normalize_numbering_file(target))
            flat = [x for s in _numids_per_table(target) for x in s]
            self.assertEqual(len(flat), len(set(flat)), '表間で numId 重複')
            # 冪等性: 2 回目は変更なし
            self.assertFalse(docx_writer.normalize_numbering_file(target))


# No 列を heading スタイルの自動採番で実装している凡例(直接 numPr ではない)。
SVC_TEMPLATE = (Path(__file__).resolve().parents[6] / 'docs' / 'specs'
                / '詳細設計_書式凡例' / 'イベント定義書' / 'サービス定義書'
                / '_サービス定義書_凡例.docx')


@unittest.skipUnless(HAS_PYTHON_DOCX, 'python-docx なし')
@unittest.skipUnless(SVC_TEMPLATE.is_file(), 'サービス定義書凡例なし')
class StyleNumberingRemediationTest(unittest.TestCase):
    """書式(スタイル)由来の自動採番も表ごとに独立させる(直接 numPr を被せる)。"""

    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

    def _ordered_styles_and_numids(self, doc):
        style_num = docx_reader._style_numbering_from_styles(doc.styles.element)
        resolver = docx_reader._numbering_resolver_from_root(
            doc.part.numbering_part.element, style_num)
        sids, numids = set(), set()
        for sid, (nid, ilvl) in style_num.items():
            fmt = resolver.fmt(nid, ilvl)
            if docx_reader._is_ordered(fmt):
                sids.add(sid)
                numids.add(nid)
        return sids, numids

    def _style_numbered_paras(self, tbl, ordered_styles):
        W = self.W
        for p in tbl.iter(W + 'p'):
            p_pr = p.find(W + 'pPr')
            if p_pr is None:
                continue
            ps = p_pr.find(W + 'pStyle')
            if ps is not None and ps.get(W + 'val') in ordered_styles:
                yield p, p_pr

    def test_style_numbering_split_per_table(self):
        import copy
        from docx import Document
        W = self.W
        doc = Document(str(SVC_TEMPLATE))
        body = doc.element.body
        ordered_styles, orig_numids = self._ordered_styles_and_numids(doc)
        self.assertTrue(ordered_styles, '凡例にスタイル由来の連番が無い')

        style_tbl = next(
            (t for t in body.findall(W + 'tbl')
             if any(self._style_numbered_paras(t, ordered_styles))), None)
        self.assertIsNotNone(style_tbl, 'スタイル採番の No 列を持つ表が無い')
        # 複製して 2 表がスタイル採番(同一 numId)を共有する状態を作る
        clone = copy.deepcopy(style_tbl)
        sect = body.find(W + 'sectPr')
        (sect.addprevious if sect is not None else body.append)(clone)

        created = docx_writer.normalize_table_numbering(doc)
        self.assertGreaterEqual(created, 1)

        # 元表: スタイルのまま据え置き(直接 numPr は付かない)
        for _, p_pr in self._style_numbered_paras(style_tbl, ordered_styles):
            self.assertIsNone(p_pr.find(W + 'numPr'), '元表に直接 numPr が付いた')
        # 複製表: 直接 numPr が被さり、元のスタイル numId とは別番号で独立採番
        clone_numids = set()
        for _, p_pr in self._style_numbered_paras(clone, ordered_styles):
            num_pr = p_pr.find(W + 'numPr')
            self.assertIsNotNone(num_pr, '複製表のスタイル採番に直接 numPr が無い')
            clone_numids.add(num_pr.find(W + 'numId').get(W + 'val'))
        self.assertTrue(clone_numids)
        self.assertTrue(clone_numids.isdisjoint(orig_numids),
                        f'複製表が元の numId を共有: {clone_numids} & {orig_numids}')
        # 冪等性
        self.assertEqual(docx_writer.normalize_table_numbering(doc), 0)


if __name__ == '__main__':
    unittest.main(verbosity=2)
