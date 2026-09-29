"""テスト仕様書xlsxの全セルをテキストへダンプする（調査用・読み取り専用）。"""
import sys

import openpyxl

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True)
with open(dst, "w", encoding="utf-8") as out:
    for ws in wb.worksheets:
        out.write(f"=== SHEET: {ws.title} (max_row={ws.max_row}, max_col={ws.max_column}) ===\n")
        for row in ws.iter_rows():
            vals = [(c.coordinate, str(c.value).replace("\n", "⏎")) for c in row if c.value is not None]
            if vals:
                out.write(" | ".join(f"{k}:{v}" for k, v in vals) + "\n")
print("dumped")
