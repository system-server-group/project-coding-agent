"""doctool xlsx リーダーのテスト (要 openpyxl)。

実行: PYTHONUTF8=1 python .claude/skills/doctool/scripts/doctool/tests/test_xlsx_reader.py
"""
import datetime
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import preview

try:
    import openpyxl
    import xlsx_reader
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

SAMPLE_DIR = Path(__file__).resolve().parents[1] / '方眼紙'


def build_xlsx(populate):
    wb = openpyxl.Workbook()
    populate(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@unittest.skipUnless(HAS_OPENPYXL, 'openpyxl なし')
class BoundaryCompressionTest(unittest.TestCase):
    """方眼紙の境界圧縮: 使われている行/列境界だけが論理グリッドになる。"""

    @classmethod
    def setUpClass(cls):
        def populate(ws):
            # 方眼紙風: 極小列に値・結合を散らす
            ws.merge_cells('B2:H2')
            ws['B2'] = 'タイトル'
            ws['B10'] = 'ラベル'
            ws.merge_cells('E10:H11')
            ws['E10'] = '値'
            ws['B12'] = datetime.datetime(2026, 6, 10)
        cls.ir = xlsx_reader.read_xlsx(build_xlsx(populate), source_name='fixture.xlsx')
        cls.table = cls.ir['sheets'][0]['blocks'][0]
        cls.cells = {(c['row'], c['col']): c for c in cls.table['cells']}

    def test_compressed_grid(self):
        # 行境界 {2,3,10,11,12,13} → 5論理行(空白帯3..9行は1論理行に縮約)
        # 列境界 {2,3,5,9} → 3論理列
        self.assertEqual(self.table['rowCount'], 5)
        self.assertEqual(self.table['colCount'], 3)

    def test_merged_cell_spans(self):
        title = self.cells[(0, 0)]
        self.assertEqual(title['blocks'][0]['text'], 'タイトル')
        self.assertEqual(title['colSpan'], 3)  # B2:H2 が全列幅に圧縮される
        value = self.cells[(2, 2)]
        self.assertEqual(value['blocks'][0]['text'], '値')
        self.assertEqual(value['rowSpan'], 2)  # E10:H11 は10〜11行に跨がる

    def test_provenance_ref(self):
        self.assertEqual(self.cells[(0, 0)]['ref'], 'B2:H2')
        self.assertEqual(self.cells[(2, 0)]['ref'], 'B10')

    def test_date_normalized(self):
        self.assertEqual(self.cells[(4, 0)]['blocks'][0]['text'], '2026-06-10')

    def test_empty_sheet_has_no_table(self):
        ir = xlsx_reader.read_xlsx(build_xlsx(lambda ws: None),
                                   source_name='empty.xlsx')
        self.assertEqual(ir['sheets'][0]['blocks'], [])


@unittest.skipUnless(HAS_OPENPYXL, 'openpyxl なし')
class FormatOverlayTest(unittest.TestCase):
    """書式オーバーレイ(with_format): 既定値以外のセル書式だけを付加する。"""

    @staticmethod
    def _populate(ws):
        from openpyxl.styles import Alignment, Font, PatternFill
        ws['B2'] = '見出し'
        ws['B2'].font = Font(bold=True, color='FF0000')
        ws['B2'].fill = PatternFill('solid', fgColor='D9D9D9')
        ws['B2'].alignment = Alignment(horizontal='center')
        ws['B3'] = '通常'

    def test_format_overlay(self):
        ir = xlsx_reader.read_xlsx(build_xlsx(self._populate),
                                   source_name='f.xlsx', with_format=True)
        cells = {(c['row'], c['col']): c
                 for c in ir['sheets'][0]['blocks'][0]['cells']}
        self.assertEqual(cells[(0, 0)]['format'], {
            'bold': True, 'color': 'FF0000', 'fill': 'D9D9D9',
            'align': 'center'})
        # 既定書式のセルには format を付けない
        self.assertNotIn('format', cells[(1, 0)])

    def test_off_by_default(self):
        ir = xlsx_reader.read_xlsx(build_xlsx(self._populate),
                                   source_name='f.xlsx')
        for c in ir['sheets'][0]['blocks'][0]['cells']:
            self.assertNotIn('format', c)


@unittest.skipUnless(HAS_OPENPYXL, 'openpyxl なし')
@unittest.skipUnless(SAMPLE_DIR.is_dir(), '方眼紙サンプルなし')
class SampleIntegrationTest(unittest.TestCase):
    def test_transition_diagram_graph(self):
        ir = xlsx_reader.read_xlsx(
            SAMPLE_DIR / '03_商品マスタ管理_画面遷移図.xlsx')
        blocks = ir['sheets'][0]['blocks']
        shapes = [b for b in blocks if b['type'] == 'shape']
        edges = [b for b in blocks if b['type'] == 'connector']
        texts = {s.get('text', '') for s in shapes}
        self.assertIn('商品登録画面', {t.replace('\n', '').replace(' ', '')
                                       for t in texts})
        self.assertGreaterEqual(len(edges), 7)
        # コネクタのID参照が図形テキストへ解決されている
        resolved = [e for e in edges if e.get('from') and e.get('to')]
        self.assertGreaterEqual(len(resolved), 7)

    def test_gridsheet_header_hierarchy(self):
        ir = xlsx_reader.read_xlsx(
            SAMPLE_DIR / '03_商品マスタ管理_画面仕様.xlsx')
        sheet = next(s for s in ir['sheets'] if s['name'] == '登録画面')
        table = sheet['blocks'][0]
        # 106 Excel列が大幅に圧縮される
        self.assertLess(table['colCount'], 40)
        texts = {c['blocks'][0]['text'] for c in table['cells']}
        self.assertIn('画面仕様', texts)
        self.assertIn('commodityName', texts)
        # 出所(元のExcel範囲)を保持
        self.assertTrue(all('ref' in c for c in table['cells']))

    def test_all_samples_readable(self):
        files = sorted(SAMPLE_DIR.glob('*.xlsx'))
        self.assertGreaterEqual(len(files), 7)
        for f in files:
            ir = xlsx_reader.read_xlsx(f, source_name=f.name)
            self.assertTrue(ir['sheets'], f'{f.name}: シートなし')
            md = preview.to_markdown(ir)
            self.assertIn('## シート:', md)


if __name__ == '__main__':
    unittest.main(verbosity=2)
