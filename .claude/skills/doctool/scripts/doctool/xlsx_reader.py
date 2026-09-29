"""doctool xlsx リーダー — Excel(方眼紙含む)を IR へ変換する (要 openpyxl)。

方眼紙の正規化(境界圧縮):
  値を持つセル・結合範囲が使う行/列の境界だけを論理グリッドの境界として採用する。
  これにより極小列の連なりや空白帯は縮約され、結合セルは rowSpan/colSpan を持つ
  論理セルになる(docx の表と同じ語彙)。各セルは出所として元の Excel 範囲 `ref` を持つ。

図形レイヤ:
  シートに紐づく drawing から、図形(テキスト+セルアンカー)・コネクタ(接続元/先の
  図形ID参照を解決した遷移エッジ)・画像を抽出する。
"""
import datetime
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter

from ir import IR_FORMAT, IR_VERSION, prune

SS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
XDR = '{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
PKG_REL = '{http://schemas.openxmlformats.org/package/2006/relationships}'


# ---------------------------------------------------------------------------
# セル層(境界圧縮)
# ---------------------------------------------------------------------------

def _cell_text(value):
    if isinstance(value, datetime.datetime):
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=' ')
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value)


def _excel_ref(r1, c1, r2, c2):
    start = f'{get_column_letter(c1)}{r1}'
    if (r1, c1) == (r2, c2):
        return start
    return f'{start}:{get_column_letter(c2)}{r2}'


def _rgb(color):
    """openpyxl の Color → 6桁 RGB(テーマ色・自動は None)。"""
    if color is None or getattr(color, 'type', None) != 'rgb':
        return None
    rgb = color.rgb
    if not isinstance(rgb, str):
        return None
    return rgb[-6:]


def _cell_format(cell, default_font):
    """セルの書式 dict(レビューに有用なものだけ。既定値は出さない)。"""
    fmt = {}
    f = cell.font
    if f.bold:
        fmt['bold'] = True
    if f.italic:
        fmt['italic'] = True
    if f.underline and f.underline != 'none':
        fmt['underline'] = True
    if f.strike:
        fmt['strike'] = True
    color = _rgb(f.color)
    if color:
        fmt['color'] = color
    if f.size and (default_font is None or f.size != default_font.size):
        size = float(f.size)
        fmt['sizePt'] = int(size) if size.is_integer() else size
    if f.name and (default_font is None or f.name != default_font.name):
        fmt['font'] = f.name
    if cell.fill is not None and cell.fill.patternType == 'solid':
        fill = _rgb(cell.fill.fgColor)
        if fill:
            fmt['fill'] = fill
    if cell.alignment is not None and cell.alignment.horizontal:
        fmt['align'] = cell.alignment.horizontal
    return fmt


def _default_font(ws):
    """ブック既定フォント(Normal スタイル)。フォント名・サイズのノイズ抑制用。"""
    try:
        return ws.parent._named_styles['Normal'].font
    except Exception:
        return None


def _sheet_table(ws, with_format=False):
    """シートのセル層を境界圧縮した table ブロックにする(値が無ければ None)。"""
    merge_origin = {(rng.min_row, rng.min_col): rng for rng in ws.merged_cells.ranges}
    default_font = _default_font(ws) if with_format else None
    items = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = _cell_text(cell.value)
            if not text.strip():
                continue
            fmt = _cell_format(cell, default_font) if with_format else None
            rng = merge_origin.get((cell.row, cell.column))
            if rng is not None:
                items.append((rng.min_row, rng.min_col,
                              rng.max_row, rng.max_col, text, fmt))
            else:
                items.append((cell.row, cell.column,
                              cell.row, cell.column, text, fmt))
    if not items:
        return None

    row_bounds = sorted({b for it in items for b in (it[0], it[2] + 1)})
    col_bounds = sorted({b for it in items for b in (it[1], it[3] + 1)})
    row_index = {v: i for i, v in enumerate(row_bounds)}
    col_index = {v: i for i, v in enumerate(col_bounds)}

    cells = []
    for r1, c1, r2, c2, text, fmt in sorted(items, key=lambda it: it[:4]):
        cell = {
            'row': row_index[r1], 'col': col_index[c1],
            'rowSpan': row_index[r2 + 1] - row_index[r1],
            'colSpan': col_index[c2 + 1] - col_index[c1],
            'ref': _excel_ref(r1, c1, r2, c2),
            'blocks': [{'type': 'paragraph', 'text': text}],
        }
        if fmt:
            cell['format'] = fmt
        cells.append(cell)
    return {'type': 'table',
            'rowCount': len(row_bounds) - 1, 'colCount': len(col_bounds) - 1,
            'cells': cells}


# ---------------------------------------------------------------------------
# 図形層(drawing)
# ---------------------------------------------------------------------------

def _resolve_target(base_part, target):
    """rels の Target を package パスへ正規化する。"""
    if target.startswith('/'):
        return target[1:]
    parts = base_part.split('/')[:-1]
    for seg in target.split('/'):
        if seg == '..':
            parts.pop()
        else:
            parts.append(seg)
    return '/'.join(parts)


def _read_rels(z, part):
    head, _, tail = part.rpartition('/')
    try:
        root = ET.fromstring(z.read(f'{head}/_rels/{tail}.rels'))
    except KeyError:
        return {}
    return {rel.get('Id'): _resolve_target(part, rel.get('Target'))
            for rel in root.findall(PKG_REL + 'Relationship')}


def _sheet_parts(z):
    """シート名 → worksheet part パス。"""
    wb_root = ET.fromstring(z.read('xl/workbook.xml'))
    rels = _read_rels(z, 'xl/workbook.xml')
    out = {}
    for sheet in wb_root.find(SS + 'sheets'):
        target = rels.get(sheet.get(R + 'id'))
        if target:
            out[sheet.get('name')] = target
    return out


def _anchor_of(anchor_el):
    pos = {}
    for key in ('from', 'to'):
        el = anchor_el.find(XDR + key)
        if el is not None:
            pos[key] = [int(el.find(XDR + 'row').text), int(el.find(XDR + 'col').text)]
    return pos or None


def _shape_text(sp):
    paras = []
    for p in sp.iter(A + 'p'):
        text = ''.join(t.text or '' for t in p.iter(A + 't'))
        if text.strip():
            paras.append(text)
    return '\n'.join(paras)


def _drawing_blocks(z, drawing_part):
    root = ET.fromstring(z.read(drawing_part))
    shapes, connectors, images = [], [], []
    id_to_text = {}
    for anchor in root:
        pos = _anchor_of(anchor)
        for sp in anchor.iter(XDR + 'sp'):
            nv = sp.find(XDR + 'nvSpPr/' + XDR + 'cNvPr')
            geom = sp.find('.//' + A + 'prstGeom')
            shape_id = int(nv.get('id')) if nv is not None else None
            text = _shape_text(sp)
            shapes.append(prune({
                'type': 'shape', 'id': shape_id,
                'name': nv.get('name') if nv is not None else None,
                'shape': geom.get('prst') if geom is not None else None,
                'anchor': pos, 'text': text,
            }))
            if shape_id is not None:
                id_to_text[shape_id] = text or (nv.get('name') if nv is not None else '')
        for cxn in anchor.iter(XDR + 'cxnSp'):
            start = cxn.find('.//' + A + 'stCxn')
            end = cxn.find('.//' + A + 'endCxn')
            connectors.append(prune({
                'type': 'connector',
                'fromId': int(start.get('id')) if start is not None else None,
                'toId': int(end.get('id')) if end is not None else None,
                'anchor': pos,
            }))
        for pic in anchor.iter(XDR + 'pic'):
            nv = pic.find(XDR + 'nvPicPr/' + XDR + 'cNvPr')
            images.append(prune({
                'type': 'image',
                'name': nv.get('name') if nv is not None else None,
                'anchor': pos,
            }))
    # コネクタの参照IDを図形テキストへ解決する(意図=遷移グラフ)
    for c in connectors:
        if 'fromId' in c:
            c['from'] = id_to_text.get(c['fromId'])
        if 'toId' in c:
            c['to'] = id_to_text.get(c['toId'])
    return shapes + connectors + images


def _sheet_drawings(z, sheet_part):
    root = ET.fromstring(z.read(sheet_part))
    rels = _read_rels(z, sheet_part)
    blocks = []
    for drawing in root.findall(SS + 'drawing'):
        target = rels.get(drawing.get(R + 'id'))
        if target and target in z.namelist():
            blocks.extend(_drawing_blocks(z, target))
    return blocks


# ---------------------------------------------------------------------------
# エントリポイント
# ---------------------------------------------------------------------------

def read_xlsx(src, source_name=None, with_format=False):
    """xlsx(パスまたはファイルオブジェクト)を IR(dict) に変換する。

    with_format=True でセルの書式(format キー: bold/italic/underline/strike/
    color/sizePt/font/fill/align)を付加する。フォント名・サイズはブック既定
    フォントと異なる場合のみ出力する。テーマ色は対象外(RGB 指定のみ)。
    """
    wb = openpyxl.load_workbook(src, data_only=True)
    if hasattr(src, 'seek'):
        src.seek(0)
    sheets = []
    with zipfile.ZipFile(src) as z:
        sheet_parts = _sheet_parts(z)
        for ws in wb.worksheets:
            blocks = []
            table = _sheet_table(ws, with_format)
            if table:
                blocks.append(table)
            part = sheet_parts.get(ws.title)
            if part:
                blocks.extend(_sheet_drawings(z, part))
            entry = prune({
                'name': ws.title,
                'hidden': True if ws.sheet_state != 'visible' else None,
            })
            entry['blocks'] = blocks
            sheets.append(entry)
    return {
        'format': IR_FORMAT,
        'version': IR_VERSION,
        'source': {'path': source_name or str(src), 'kind': 'xlsx'},
        'sheets': sheets,
    }
