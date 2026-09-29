"""doctool xlsx fitrows のテスト (要 lxml / openpyxl)。

実行: PYTHONUTF8=1 python .claude/skills/doctool/scripts/doctool/tests/test_xlsx_fitrows.py

要点:
  - 「Excel で開いて保存」は、行タグへの ht 属性注入で模擬する(customHeight は付けない)。
  - round0: 固定高(customHeight)のケース行の高さ退避と固定解除 → 終了コード3。
  - round0b: 退避行の実測回収 → 足りる行は元の高さへ復元、足りない行は引き上げ、
    頭打ちの行は自動高さのまま断片測定へ送る。
  - 1回目: 頭打ち行(ht>=409.4)の検出とスクラッチ(測定用断片)の生成 → 終了コード3。
  - 2回目: 断片実測の合算 → 行挿入+縦マージ+明示行高の適用、スクラッチ削除、印刷範囲追従。
  - 保存忘れ・内容変更・断片頭打ち・ページ高超過の各エラー経路。
"""
import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

try:
    from lxml import etree
    import openpyxl
    from openpyxl.styles import Alignment
    import xlsx_fitrows
    import xlsx_writer
    HAS_DEPS = True
except ImportError:
    HAS_DEPS = False

SS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def _fixture(long_lines=40, as_block=False, fixed_ht=None):
    """ケース行 6..8(1-1/1-2/1-3)を持つ検証用ワークブックを作る。

    1-2(行7)の B 列に long_lines 行の長文を置く。as_block=True のときは
    1-2 を「分割済みブロック」(行7-8 を縦マージ・固定高)にし、1-3 は行9になる。
    fixed_ht を渡すと各ケース行に固定高(customHeight)を与える(round0 の対象になる)。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'テストケース'
    wrap = Alignment(wrap_text=True, vertical='top')
    for i, h in enumerate(['No', '内容', '期待']):
        ws.cell(row=5, column=1 + i, value=h)
    long_text = '\n'.join(f'手順{i:02d}' for i in range(1, long_lines + 1))
    rows = [('1-1', '短い', 'r1'), ('1-2', long_text, 'r2'), ('1-3', '普通\n2行', 'r3')]
    r = 6
    case_rows = []
    for no, naiyo, kitai in rows:
        for k, v in enumerate((no, naiyo, kitai)):
            ws.cell(row=r, column=1 + k, value=v).alignment = wrap
        case_rows.append(r)
        if as_block and no == '1-2':
            for col in 'ABC':
                ws.merge_cells(f'{col}7:{col}8')
            ws.row_dimensions[7].height = 300
            ws.row_dimensions[8].height = 300
            r += 1                     # ブロックぶん1行余計に進める
        r += 1
    if fixed_ht:
        for cr in case_rows:
            ws.row_dimensions[cr].height = fixed_ht
    ws.print_area = f'A1:C{r + 1}'
    ws.print_title_rows = '1:5'
    d = Path(tempfile.mkdtemp())
    path = d / 'target.xlsx'
    wb.save(path)
    return path


def _zip_parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def _inject_hts(path, hts):
    """Excel 保存の模擬: 指定行に ht を付与する(customHeight は付けない)。"""
    parts = _zip_parts(path)
    xml = parts['xl/worksheets/sheet1.xml'].decode('utf-8')
    for r, ht in hts.items():
        m = re.search(r'<row r="%d"[^>]*?>' % r, xml)
        if m is None:
            raise AssertionError(f'行 {r} が見つかりません')
        tag = m.group(0)
        new = re.sub(r' ht="[0-9.]+"', '', tag)[:-1] + f' ht="{ht}">'
        xml = xml.replace(tag, new, 1)
    parts['xl/worksheets/sheet1.xml'] = xml.encode('utf-8')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)


def _rows(path):
    """{行番号: (ht, customHeight)}"""
    root = etree.fromstring(_zip_parts(path)['xl/worksheets/sheet1.xml'])
    return {int(row.get('r')): (row.get('ht'), row.get('customHeight'))
            for row in root.iter(f'{{{SS}}}row')}


def _merges(path):
    root = etree.fromstring(_zip_parts(path)['xl/worksheets/sheet1.xml'])
    mc = root.find(f'{{{SS}}}mergeCells')
    return {m.get('ref') for m in mc.findall(f'{{{SS}}}mergeCell')} if mc is not None else set()


def _print_area(path):
    root = etree.fromstring(_zip_parts(path)['xl/workbook.xml'])
    dnc = root.find(f'{{{SS}}}definedNames')
    for dn in (dnc if dnc is not None else []):
        if dn.get('name') == '_xlnm.Print_Area':
            return dn.text
    return None


def _marker_rows(path, marker, key_re, key_col='D'):
    """キー列にマーカー/キーを持つ行番号のリスト(作業行の存在確認)。"""
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    out = []
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str):
                continue
            if openpyxl.utils.get_column_letter(c.column) != key_col:
                continue
            if c.value.startswith(marker) or key_re.match(c.value):
                out.append(c.row)
    return out


def _scratch_rows(path, key_col='D'):
    return _marker_rows(path, xlsx_fitrows.MARKER, xlsx_fitrows._KEY_RE, key_col)


def _backup_rows(path, key_col='D'):
    return _marker_rows(path, xlsx_fitrows.MARKER_HT, xlsx_fitrows._HT_KEY_RE, key_col)


@unittest.skipUnless(HAS_DEPS, 'lxml / openpyxl なし')
class FitrowsTest(unittest.TestCase):
    def _spec(self, path, **kw):
        spec = {'file': str(path), 'anchorColumn': 'A', 'anchorPattern': r'^\d+-\d+$',
                'columns': 'A:C', 'maxRowPt': 400, 'pageHeightPt': 2000}
        spec.update(kw)
        return spec

    def _measure(self, path, marker_ht=48.0, frag_ht=24.0):
        """スクラッチ全行へ実測 ht を注入する(Excel 保存の模擬)。"""
        rows = _scratch_rows(path)
        _inject_hts(path, {r: (marker_ht if i == 0 else frag_ht)
                           for i, r in enumerate(sorted(rows))})
        return len(rows) - 1        # 断片行数

    def _measure_backup(self, path, case_hts, marker_ht=48.0):
        """退避マーカーへ ht を付け(=保存済みの印)、ケース行へ実測 ht を注入する。"""
        rows = sorted(_backup_rows(path))
        _inject_hts(path, {rows[0]: marker_ht, **case_hts})

    def test_no_capped_no_block_is_noop(self):
        path = _fixture()
        _inject_hts(path, {6: 24, 7: 100, 8: 48})
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 0)
        self.assertEqual(_scratch_rows(path), [])

    def test_round1_writes_scratch_and_requests_save(self):
        path = _fixture(long_lines=40)
        _inject_hts(path, {6: 24, 7: 409.5, 8: 48})
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)
        scratch = _scratch_rows(path)
        # マーカー1行 + 断片(A:1 + B:40 + C:1)= 43 行
        self.assertEqual(len(scratch), 43)
        # スクラッチ行は自動高さ(ht なし)で作られる
        rows = _rows(path)
        self.assertTrue(all(rows[r] == (None, None) for r in scratch))
        # 印刷範囲はスクラッチまで広がらない
        self.assertEqual(_print_area(path), "'テストケース'!$A$1:$C$10")

    def test_round2_without_save_detected(self):
        path = _fixture(long_lines=40)
        _inject_hts(path, {6: 24, 7: 409.5, 8: 48})
        xlsx_fitrows.run(self._spec(path))
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)   # 保存忘れ
        self.assertTrue(_scratch_rows(path))                      # スクラッチは維持

    def test_full_round_trip_splits_capped_row(self):
        path = _fixture(long_lines=40)
        _inject_hts(path, {6: 24, 7: 409.5, 8: 48})
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)
        self._measure(path)                    # B列 40断片 × 24pt = 必要高960pt
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 0)
        # 960+3 → 3行 × ceil(963/3*4)/4 = 321pt
        rows = _rows(path)
        for r in (7, 8, 9):
            self.assertEqual(rows[r], ('321', '1'))
        merges = _merges(path)
        for col in 'ABC':
            self.assertIn(f'{col}7:{col}9', merges)
        ws = openpyxl.load_workbook(path).active
        self.assertEqual(ws['A10'].value, '1-3')                  # 下のケースは +2 追従
        self.assertEqual(ws['B7'].value.count('\n'), 39)          # 全文は先頭セルに残る
        self.assertEqual(_print_area(path), "'テストケース'!$A$1:$C$12")
        self.assertEqual(_scratch_rows(path), [])                 # スクラッチは削除
        # 通常行(1-1)は触らない(自動高さのまま)
        self.assertEqual(rows[6], ('24', None))
        # 新規作成された mergeCells はスキーマ順の位置に入る(末尾だと Excel が開けない)
        xml = _zip_parts(path)['xl/worksheets/sheet1.xml'].decode('utf-8')
        self.assertLess(xml.index('<mergeCells'), xml.index('<pageMargins'))

    def test_block_sufficient_is_noop_after_measure(self):
        path = _fixture(long_lines=10, as_block=True)             # 必要高240 < 600
        _inject_hts(path, {6: 24, 9: 48})
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)   # ブロックは常に再測定
        self._measure(path)
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 0)
        rows = _rows(path)
        self.assertEqual(float(rows[7][0]), 300.0)                # 現状維持(余裕は保つ)
        self.assertEqual(rows[7][1], '1')
        self.assertEqual(_scratch_rows(path), [])

    def test_block_extended_when_rows_insufficient(self):
        path = _fixture(long_lines=54, as_block=True)             # 必要高1296 > 400*2
        _inject_hts(path, {6: 24, 9: 48})
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)
        self._measure(path)                    # 1296+3 → 4行 × 1299/4 = 324.75pt
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 0)
        rows = _rows(path)
        for r in (7, 8, 9, 10):
            self.assertEqual(rows[r], ('324.75', '1'))
        merges = _merges(path)
        for col in 'ABC':
            self.assertIn(f'{col}7:{col}10', merges)
            self.assertNotIn(f'{col}7:{col}8', merges)            # 旧結合は除去
        ws = openpyxl.load_workbook(path).active
        self.assertEqual(ws['A11'].value, '1-3')                  # 2行追加ぶん追従

    def test_page_height_guard_blocks_apply(self):
        path = _fixture(long_lines=40)
        _inject_hts(path, {6: 24, 7: 409.5, 8: 48})
        xlsx_fitrows.run(self._spec(path))
        self._measure(path)
        rc = xlsx_fitrows.run(self._spec(path, pageHeightPt=500))  # 960pt は収まらない
        self.assertEqual(rc, 1)
        self.assertNotIn('B7:B9', _merges(path))                  # 分割は適用されない
        self.assertEqual(_scratch_rows(path), [])                 # スクラッチは掃除される

    def test_capped_fragment_is_error(self):
        path = _fixture(long_lines=40)
        _inject_hts(path, {6: 24, 7: 409.5, 8: 48})
        xlsx_fitrows.run(self._spec(path))
        scratch = sorted(_scratch_rows(path))
        _inject_hts(path, {r: (48.0 if i == 0 else 24.0)
                           for i, r in enumerate(scratch)})
        _inject_hts(path, {scratch[2]: 409.5})                    # 断片単体が頭打ち
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 1)
        self.assertEqual(_scratch_rows(path), [])

    # -- round0 / round0b: 固定高(customHeight)行の再測定 ------------------

    def test_round0_backs_up_and_clears_fixed_height(self):
        path = _fixture(long_lines=3, fixed_ht=100)
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)
        rows = _rows(path)
        for r in (6, 7, 8):                                   # 固定が外れて自動高さに
            self.assertEqual(rows[r], (None, None))
        backup = sorted(_backup_rows(path))
        self.assertEqual(len(backup), 4)                      # マーカー1 + 退避3
        ws = openpyxl.load_workbook(path).active
        self.assertEqual([ws.cell(row=r, column=4).value for r in backup[1:]],
                         ['ht|6|100', 'ht|7|100', 'ht|8|100'])
        self.assertEqual(_scratch_rows(path), [])             # 断片スクラッチはまだ無い

    def test_fixed_height_raised_when_measured_taller(self):
        path = _fixture(long_lines=3, fixed_ht=100)
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)
        self._measure_backup(path, {7: 150.0})                # 1-2 の実測が固定高を超える
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 0)
        rows = _rows(path)
        self.assertEqual(rows[7], ('153', '1'))               # 150 + slackPt(3)
        for r in (6, 8):                                      # 足りる行は元の高さへ復元
            self.assertEqual(rows[r], ('100', '1'))
        self.assertEqual(_backup_rows(path), [])              # 退避行は削除

    def test_fixed_height_restored_when_sufficient(self):
        path = _fixture(long_lines=3, fixed_ht=300)
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)
        self._measure_backup(path, {7: 150.0})                # 300 で足りている
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 0)
        for r in (6, 7, 8):
            self.assertEqual(_rows(path)[r], ('300', '1'))    # 意図した余裕は維持
        self.assertEqual(_backup_rows(path), [])

    def test_fixed_height_exactly_sufficient_is_not_raised(self):
        """実測がちょうど現行の固定高と同じ行は見切れていない = 触らない。"""
        path = _fixture(long_lines=3, fixed_ht=288)
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)
        self._measure_backup(path, {7: 288.0})
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 0)
        self.assertEqual(_rows(path)[7], ('288', '1'))        # slackPt を足して膨らませない

    def test_fixed_height_slack_is_configurable(self):
        path = _fixture(long_lines=3, fixed_ht=100)
        spec = self._spec(path, slackPt=24)
        self.assertEqual(xlsx_fitrows.run(spec), 3)
        self._measure_backup(path, {7: 150.0})
        self.assertEqual(xlsx_fitrows.run(spec), 0)
        self.assertEqual(_rows(path)[7], ('174', '1'))        # 150 + 24

    def test_fixed_height_over_page_height_is_error(self):
        path = _fixture(long_lines=3, fixed_ht=100)
        spec = self._spec(path, pageHeightPt=200)
        self.assertEqual(xlsx_fitrows.run(spec), 3)
        self._measure_backup(path, {7: 300.0})                # 300 は1ページに収まらない
        self.assertEqual(xlsx_fitrows.run(spec), 1)
        self.assertEqual(_rows(path)[7], ('100', '1'))        # 引き上げず元へ戻す
        self.assertEqual(_backup_rows(path), [])              # 固定解除の状態は残さない

    def test_page_height_ng_does_not_block_fragment_measure(self):
        path = _fixture(long_lines=40, fixed_ht=100)
        spec = self._spec(path, pageHeightPt=200)
        self.assertEqual(xlsx_fitrows.run(spec), 3)
        self._measure_backup(path, {6: 300.0, 7: 409.5})   # 1-1 はページ超過 / 1-2 は頭打ち
        self.assertEqual(xlsx_fitrows.run(spec), 3)        # NG があっても断片測定へ進む
        self.assertEqual(_rows(path)[6], ('100', '1'))     # NG 行は元へ戻す
        self.assertTrue(_scratch_rows(path))

    def test_round0b_without_save_keeps_backup(self):
        path = _fixture(long_lines=3, fixed_ht=100)
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)   # 保存忘れ
        self.assertEqual(len(_backup_rows(path)), 4)              # 退避は維持
        self.assertEqual(_rows(path)[7], (None, None))

    def test_fixed_height_capped_goes_to_fragment_measure(self):
        path = _fixture(long_lines=40, fixed_ht=100)
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)   # round0
        self._measure_backup(path, {7: 409.5})                    # 実測が頭打ち
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 3)   # round0b → round1
        self.assertEqual(_backup_rows(path), [])                  # 退避は片付く
        self.assertEqual(len(_scratch_rows(path)), 43)            # 断片測定へ進んだ
        rows = _rows(path)
        self.assertEqual(rows[7], ('409.5', None))                # 頭打ち行は自動高さのまま
        for r in (6, 8):
            self.assertEqual(rows[r], ('100', '1'))               # 他は復元済み
        self._measure(path)                    # B列 40断片 × 24pt = 必要高960pt
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 0)   # round2
        rows = _rows(path)
        for r in (7, 8, 9):
            self.assertEqual(rows[r], ('321', '1'))               # 960+3 → 3行 × 321pt
        for col in 'ABC':
            self.assertIn(f'{col}7:{col}9', _merges(path))
        self.assertEqual(openpyxl.load_workbook(path).active['A10'].value, '1-3')

    def test_content_change_after_scratch_is_error(self):
        path = _fixture(long_lines=40)
        _inject_hts(path, {6: 24, 7: 409.5, 8: 48})
        xlsx_fitrows.run(self._spec(path))
        wb = openpyxl.load_workbook(path)                          # 測定後にセルを書き換え
        wb.active['B7'] = '書き換えた\n2行'
        wb.save(path)
        self._measure(path)
        self.assertEqual(xlsx_fitrows.run(self._spec(path)), 1)
        self.assertEqual(_scratch_rows(path), [])


if __name__ == '__main__':
    unittest.main(verbosity=2)
