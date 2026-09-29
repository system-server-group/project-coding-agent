"""doctool xlsx ライターのテスト (要 lxml / openpyxl)。

実行: PYTHONUTF8=1 python .claude/skills/doctool/scripts/doctool/tests/test_xlsx_writer.py

要点:
  - worksheet / styles / drawing のパートだけを編集し、他パート(画像等)を byte 単位で保持する。
  - 行の挿入/削除で 行・結合・**描画アンカー**が正しく追従する。
  - openpyxl の load→save はモデル化しないパートを落とすことを対比で示す。
"""
import base64
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

try:
    from lxml import etree
    import openpyxl
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font
    import xlsx_writer
    HAS_DEPS = True
except ImportError:
    HAS_DEPS = False

XDR = 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
# 1x1 透過 PNG
PNG_1PX = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8z8BQDwAE'
    'hQGAhKmMIQAAAABJRU5ErkJggg==')


def _template_with_image(image_anchor='A10'):
    """ヘッダー枠＋見本データ行(row6)＋A10 にアンカーした画像を持つ雛形を作る。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'テストケース'
    ws.merge_cells('A1:B1')
    ws['A1'] = 'システム名'
    ws['C1'] = 'サンプルシステム'
    for i, h in enumerate(['No', '系', '観点', '前提', '操作', '期待']):
        ws.cell(row=5, column=1 + i, value=h).font = Font(bold=True)
    for i, v in enumerate(['1', '正常', '観点', '前提', '1. 開く\n2. 押す', '結果']):
        ws.cell(row=6, column=1 + i, value=v).alignment = Alignment(wrap_text=True)
    d = Path(tempfile.mkdtemp())
    png = d / 'img.png'
    png.write_bytes(PNG_1PX)
    img = Image(str(png))
    img.anchor = image_anchor
    ws.add_image(img)
    path = d / 'tpl.xlsx'
    wb.save(path)
    return path


def _zip_parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def _anchor_rows(path):
    """全 drawing パートの xdr:from/to の行(0 基点)を集める。"""
    rows = []
    for name, data in _zip_parts(path).items():
        if name.startswith('xl/drawings/drawing') and name.endswith('.xml'):
            root = etree.fromstring(data)
            for tag in ('from', 'to'):
                for pos in root.iter(f'{{{XDR}}}{tag}'):
                    r = pos.find(f'{{{XDR}}}row')
                    if r is not None:
                        rows.append(int(r.text))
    return rows


@unittest.skipUnless(HAS_DEPS, 'lxml / openpyxl なし')
class CellOpsTest(unittest.TestCase):
    def _run(self, ops, anchor='A10'):
        tpl = _template_with_image(anchor)
        out = tpl.with_name('out.xlsx')
        w = xlsx_writer.apply_spec({'template': str(tpl), 'output': str(out),
                                      'operations': ops})
        return out, w

    def test_set_cell_and_merge_anchor(self):
        out, _ = self._run([{'op': 'setCell', 'ref': 'B1', 'text': 'ラベル'}])
        wb = openpyxl.load_workbook(out)
        self.assertEqual(wb['テストケース']['A1'].value, 'ラベル')  # 結合起点へ

    def test_replace_text(self):
        out, _ = self._run([{'op': 'replaceText', 'ref': 'C1',
                             'find': 'サンプル', 'replace': '文書管理'}])
        self.assertEqual(openpyxl.load_workbook(out)['テストケース']['C1'].value,
                         '文書管理システム')

    def test_clear_cell(self):
        out, _ = self._run([{'op': 'clearCell', 'ref': 'C1'}])
        self.assertIsNone(openpyxl.load_workbook(out)['テストケース']['C1'].value)

    def test_image_preserved(self):
        out, _ = self._run([{'op': 'setCell', 'ref': 'C1', 'text': 'x'}])
        parts = _zip_parts(out)
        self.assertIn('xl/media/image1.png', parts)
        self.assertEqual(parts['xl/media/image1.png'], PNG_1PX)


@unittest.skipUnless(HAS_DEPS, 'lxml / openpyxl なし')
class MergeCellsPositionTest(unittest.TestCase):
    def test_new_mergecells_inserted_in_schema_order(self):
        # 結合が1つも無いシートへの新規作成。worksheet 末尾に追加すると
        # pageMargins より後ろになり、スキーマ違反で Excel がファイルを開けない
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'x'
        ws['A2'] = 'y'
        d = Path(tempfile.mkdtemp())
        tpl = d / 'tpl.xlsx'
        wb.save(tpl)
        out = d / 'out.xlsx'
        xlsx_writer.apply_spec({'template': str(tpl), 'output': str(out),
                                'operations': [{'op': 'mergeCells', 'ref': 'A1:A2'}]})
        xml = _zip_parts(out)['xl/worksheets/sheet1.xml'].decode('utf-8')
        self.assertIn('<mergeCells', xml)
        self.assertLess(xml.index('<mergeCells'), xml.index('<pageMargins'))


@unittest.skipUnless(HAS_DEPS, 'lxml / openpyxl なし')
class SetFormatTest(unittest.TestCase):
    def test_font_fill_align(self):
        tpl = _template_with_image()
        out = tpl.with_name('out.xlsx')
        xlsx_writer.apply_spec({'template': str(tpl), 'output': str(out),
                                  'operations': [{'op': 'setFormat', 'ref': 'C1',
                                   'format': {'bold': True, 'color': 'FF0000',
                                              'fill': 'FFFF00', 'align': 'center'}}]})
        c = openpyxl.load_workbook(out)['テストケース']['C1']
        self.assertTrue(c.font.bold)
        self.assertTrue(str(c.font.color.rgb).endswith('FF0000'))
        self.assertEqual(c.fill.fgColor.rgb[-6:], 'FFFF00')
        self.assertEqual(c.alignment.horizontal, 'center')


@unittest.skipUnless(HAS_DEPS, 'lxml / openpyxl なし')
class FillTableTest(unittest.TestCase):
    ROWS = [[1, '正常', 'o1', 'p1', '1. a', 'r1'],
            [2, '異常', 'o2', 'p2', '1. b', 'r2'],
            [3, '正常', 'o3', 'p3', '1. c', 'r3']]

    def test_overwrite_default_no_shift(self):
        # 既定は上書き: 行数もアンカーも変えない(固定表領域向け)
        tpl = _template_with_image('A10')
        out = tpl.with_name('out.xlsx')
        xlsx_writer.apply_spec({'template': str(tpl), 'output': str(out),
                                  'operations': [{'op': 'fillTable', 'templateRow': 6,
                                   'rows': self.ROWS}]})
        ws = openpyxl.load_workbook(out)['テストケース']
        self.assertEqual(ws['A6'].value, '1')
        self.assertEqual(ws['C7'].value, 'o2')
        self.assertEqual(ws['A8'].value, '3')
        self.assertTrue(ws['E7'].alignment.wrap_text)   # 見本行の書式が複製される
        self.assertEqual(_anchor_rows(out), [9])        # アンカーは動かない
        self.assertEqual(_zip_parts(out)['xl/media/image1.png'], PNG_1PX)

    def test_insert_mode_shifts_anchor(self):
        # insert:true のとき 1 行上書き＋2 行挿入し、画像アンカーが +2 追従
        tpl = _template_with_image('A10')
        self.assertEqual(_anchor_rows(tpl), [9])
        out = tpl.with_name('out.xlsx')
        xlsx_writer.apply_spec({'template': str(tpl), 'output': str(out),
                                  'operations': [{'op': 'fillTable', 'templateRow': 6,
                                   'insert': True, 'rows': self.ROWS}]})
        ws = openpyxl.load_workbook(out)['テストケース']
        self.assertEqual(ws['A6'].value, '1')
        self.assertEqual(ws['A8'].value, '3')
        self.assertEqual(_anchor_rows(out), [11])
        self.assertEqual(_zip_parts(out)['xl/media/image1.png'], PNG_1PX)


@unittest.skipUnless(HAS_DEPS, 'lxml / openpyxl なし')
class DeleteRowAnchorTest(unittest.TestCase):
    def test_delete_shifts_anchor_up(self):
        tpl = _template_with_image('A10')
        out = tpl.with_name('out.xlsx')
        xlsx_writer.apply_spec({'template': str(tpl), 'output': str(out),
                                  'operations': [{'op': 'deleteRow', 'row': 5}]})
        self.assertEqual(_anchor_rows(out), [8])       # 行5削除で 9→8
        ws = openpyxl.load_workbook(out)['テストケース']
        self.assertEqual(ws['A5'].value, '1')          # 旧 row6 が row5 へ


@unittest.skipUnless(HAS_DEPS, 'lxml / openpyxl なし')
class OpenpyxlContrastTest(unittest.TestCase):
    """対比: openpyxl の load→save はモデル化しないパートを落とす。"""

    def test_openpyxl_drops_unmodeled_parts(self):
        tpl = _template_with_image()
        # モデル化されない任意パートを注入(図形/カスタムXML の代理)
        injected = tpl.with_name('injected.xlsx')
        with zipfile.ZipFile(tpl) as zin, \
                zipfile.ZipFile(injected, 'w', zipfile.ZIP_DEFLATED) as zout:
            for n in zin.namelist():
                zout.writestr(n, zin.read(n))
            zout.writestr('xl/drawings/vmlDrawing_extra.vml', b'<xml>shape</xml>')
        self.assertIn('xl/drawings/vmlDrawing_extra.vml', _zip_parts(injected))
        out = injected.with_name('opxl.xlsx')
        openpyxl.load_workbook(injected).save(out)
        self.assertNotIn('xl/drawings/vmlDrawing_extra.vml', _zip_parts(out))


def _print_area(path, name='_xlnm.Print_Area'):
    """出力の workbook.xml から Print_Area(等)の値を返す。"""
    root = etree.fromstring(_zip_parts(path)['xl/workbook.xml'])
    dnc = root.find(f'{{{xlsx_writer.SS}}}definedNames')
    for dn in (dnc if dnc is not None else []):
        if dn.get('name') == name:
            return dn.text
    return None


@unittest.skipUnless(HAS_DEPS, 'lxml / openpyxl なし')
class PrintAreaTest(unittest.TestCase):
    """記入が既存の印刷範囲をはみ出したら、範囲を記入セルまで自動拡張する。"""

    def _tpl(self):
        tpl = _template_with_image('A10')
        wb = openpyxl.load_workbook(tpl)
        ws = wb['テストケース']
        ws.print_area = 'A1:F8'            # 見本行(row6)を含む狭い印刷範囲
        ws.print_title_rows = '1:5'        # 見出し繰返し(Print_Titles): 変更されない想定
        wb.save(tpl)
        return tpl

    def _run(self, ops):
        tpl = self._tpl()
        out = tpl.with_name('out.xlsx')
        xlsx_writer.apply_spec({'template': str(tpl), 'output': str(out),
                                'operations': ops})
        return out

    def test_overwrite_expands_bottom(self):
        out = self._run([{'op': 'fillTable', 'templateRow': 6,
                          'rows': [[str(i)] for i in range(10)]}])   # row6..15
        self.assertEqual(_print_area(out), "'テストケース'!$A$1:$F$15")
        self.assertEqual(_print_area(out, '_xlnm.Print_Titles'),
                         "'テストケース'!$1:$5")                       # 見出しは不変

    def test_insert_shifts_existing_bottom(self):
        out = self._run([{'op': 'fillTable', 'templateRow': 6, 'insert': True,
                          'rows': [[str(i)] for i in range(10)]}])   # 挿入9行
        # 旧 row7,8(範囲内)が +9 され 16,17 へ。記入は row6..15。→ 下端 17
        self.assertEqual(_print_area(out), "'テストケース'!$A$1:$F$17")

    def test_delete_shifts_existing_bottom(self):
        out = self._run([{'op': 'deleteRow', 'row': 7}])            # 範囲内の1行削除
        self.assertEqual(_print_area(out), "'テストケース'!$A$1:$F$7")

    def test_inside_area_unchanged(self):
        out = self._run([{'op': 'setCell', 'ref': 'B2', 'text': 'x'}])
        self.assertEqual(_print_area(out), "'テストケース'!$A$1:$F$8")

    def test_expands_columns_and_rows(self):
        out = self._run([{'op': 'setCell', 'ref': 'K20', 'text': 'x'}])
        self.assertEqual(_print_area(out), "'テストケース'!$A$1:$K$20")

    def test_no_area_not_created(self):
        tpl = _template_with_image('A10')          # print_area 未設定
        out = tpl.with_name('out.xlsx')
        xlsx_writer.apply_spec({'template': str(tpl), 'output': str(out),
                                'operations': [{'op': 'fillTable', 'templateRow': 6,
                                                'rows': [[str(i)] for i in range(10)]}]})
        self.assertIsNone(_print_area(out))        # 未定義シートには作らない


@unittest.skipUnless(HAS_DEPS, 'lxml / openpyxl なし')
class RowHeightMergeOpsTest(unittest.TestCase):
    """setRowHeight / mergeCells の追加操作。

    画像フィクスチャ(要 Pillow)とは独立の検証のため、雛形は画像なしで作る。
    """

    @staticmethod
    def _tpl():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'テストケース'
        ws.merge_cells('A1:B1')
        ws['A1'] = 'システム名'
        for i, v in enumerate(['1', '正常', '観点']):
            ws.cell(row=6, column=1 + i, value=v)
        d = Path(tempfile.mkdtemp())
        path = d / 'tpl.xlsx'
        wb.save(path)
        return path

    def _run(self, ops):
        tpl = self._tpl()
        out = tpl.with_name('out.xlsx')
        xlsx_writer.apply_spec({'template': str(tpl), 'output': str(out),
                                'operations': ops})
        return out

    def _row_attrs(self, path, r):
        root = etree.fromstring(_zip_parts(path)['xl/worksheets/sheet1.xml'])
        for row in root.iter(f'{{{xlsx_writer.SS}}}row'):
            if row.get('r') == str(r):
                return row.get('ht'), row.get('customHeight')
        return None, None

    def test_set_row_height_fixed(self):
        out = self._run([{'op': 'setRowHeight', 'row': 6, 'heightPt': 240.5}])
        self.assertEqual(self._row_attrs(out, 6), ('240.5', '1'))

    def test_set_row_height_creates_missing_row(self):
        out = self._run([{'op': 'setRowHeight', 'rows': [6, 20], 'heightPt': 96}])
        self.assertEqual(self._row_attrs(out, 20), ('96', '1'))

    def test_set_row_height_auto_clears(self):
        tpl = self._tpl()
        wb = openpyxl.load_workbook(tpl)
        wb['テストケース'].row_dimensions[6].height = 100   # customHeight 付きの固定高
        wb.save(tpl)
        out = tpl.with_name('out.xlsx')
        xlsx_writer.apply_spec({'template': str(tpl), 'output': str(out),
                                'operations': [{'op': 'setRowHeight', 'row': 6,
                                                'auto': True}]})
        self.assertEqual(self._row_attrs(out, 6), (None, None))

    def test_set_row_height_over_max_rejected(self):
        with self.assertRaises(ValueError):
            self._run([{'op': 'setRowHeight', 'row': 6, 'heightPt': 409.6}])

    def test_set_row_height_requires_exactly_one_mode(self):
        with self.assertRaises(ValueError):
            self._run([{'op': 'setRowHeight', 'row': 6}])
        with self.assertRaises(ValueError):
            self._run([{'op': 'setRowHeight', 'row': 6, 'heightPt': 24, 'auto': True}])

    def test_merge_cells_adds_vertical(self):
        out = self._run([{'op': 'mergeCells', 'ref': 'A6:A8'}])
        ws = openpyxl.load_workbook(out)['テストケース']
        self.assertIn('A6:A8', {str(r) for r in ws.merged_cells.ranges})

    def test_merge_cells_same_ref_is_noop(self):
        out = self._run([{'op': 'mergeCells', 'ref': 'A1:B1'}])   # 雛形に既にある
        ws = openpyxl.load_workbook(out)['テストケース']
        self.assertEqual(sum(1 for r in ws.merged_cells.ranges if str(r) == 'A1:B1'), 1)

    def test_merge_cells_overlap_rejected(self):
        with self.assertRaises(ValueError):
            self._run([{'op': 'mergeCells', 'ref': 'B1:B2'}])     # A1:B1 と部分重複


if __name__ == '__main__':
    unittest.main(verbosity=2)
